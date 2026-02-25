import json
import csv
from pathlib import Path
from bs4 import BeautifulSoup

def load_html(path: str) -> BeautifulSoup:
    with open(path, "r", encoding="utf-8") as f:
        return BeautifulSoup(f.read(), "html.parser")

def load_json(path: str) -> dict:
    """Load JSON file"""
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

def load_csv(path: str) -> list[dict]:
    """Load CSV file and return as list of dicts"""
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(row)
    return rows

def load_xlsx(path: str) -> dict:
    """Load Excel file using openpyxl"""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl required for Excel support: pip install openpyxl")
    
    wb = openpyxl.load_workbook(path)
    sheets = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        headers = None
        
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i == 1:
                headers = row
            else:
                row_dict = {headers[j]: val for j, val in enumerate(row) if headers and j < len(headers)}
                rows.append(row_dict)
        
        sheets[sheet_name] = rows
    
    return sheets

def detect_file_type(path: str) -> str:
    """Detect file type from extension"""
    path_obj = Path(path)
    ext = path_obj.suffix.lower()
    
    if ext == ".html":
        return "html"
    elif ext == ".json":
        return "json"
    elif ext == ".csv":
        return "csv"
    elif ext in [".xlsx", ".xls"]:
        return "xlsx"
    else:
        return "unknown"

def load_file(path: str):
    """Universal file loader - returns appropriate format"""
    file_type = detect_file_type(path)
    
    if file_type == "html":
        return load_html(path), file_type
    elif file_type == "json":
        return load_json(path), file_type
    elif file_type == "csv":
        return load_csv(path), file_type
    elif file_type == "xlsx":
        return load_xlsx(path), file_type
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
